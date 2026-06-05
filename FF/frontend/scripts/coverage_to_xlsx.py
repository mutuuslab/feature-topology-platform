#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""기능명세서 xlsx에 '플랫폼 반영 현황' 시트를 추가한다.
- 입력: 원본 기능명세서 xlsx + src/data/spec.json(요구사항) + specCoverage 매핑(아래 COVERAGE)
- 출력: 원본을 보존한 사본에 시트 3개 추가 → '..._플랫폼반영현황.xlsx'
사용: python scripts/coverage_to_xlsx.py ["<원본 xlsx>"]
"""
import json, os, shutil, sys, re
from collections import OrderedDict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

SRC = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\TaehoKim\Downloads\Feature Flag 플랫폼_기능명세서_2026-06-01_2차revision.xlsx"
SPEC = os.path.join(os.path.dirname(__file__), "..", "src", "data", "spec.json")
OUT = os.path.join(os.path.dirname(SRC), "Feature Flag 플랫폼_기능명세서_플랫폼반영현황.xlsx")

# specCoverage.ts 와 동일한 매핑 (family → 상태·반영 화면)
COVERAGE = {
    'FR-REG':  ('완료',  ['Catalog', 'Feature Detail', 'Release Readiness', 'ChangeSet']),
    'FR-CAT':  ('완료',  ['Catalog', 'Taxonomy']),
    'FR-DEP':  ('완료',  ['Topology', 'Edge Editor', 'Impact']),
    'FR-POL':  ('완료',  ['Policy Lifecycle', 'Control Point']),
    'FR-AUTH': ('완료',  ['Runtime Sim']),
    'FR-ROL':  ('완료',  ['OTA Campaign']),
    'FR-EXP':  ('완료',  ['실험·효과검증']),
    'FR-PVER': ('완료',  ['Policy Lifecycle']),
    'FR-EXC':  ('완료',  ['예외 정책']),
    'FR-CON':  ('완료',  ['정책 충돌']),
    'FR-RBAC': ('완료',  ['Permissions']),
    'FR-AUD':  ('완료',  ['Audit Log']),
    'FR-VAR':  ('완료',  ['Variant Matrix']),
    'FR-SUP':  ('완료',  ['Supplier Portal(책임·인수)', 'API Package']),
    'FR-TGT':  ('완료',  ['Variant Matrix', 'Campaign']),
    'FR-RTE':  ('완료',  ['Runtime Sim']),
    'FR-LPC':  ('완료',  ['Runtime Sim']),
    'FR-SFD':  ('완료',  ['Kill Switch·Safe Default']),
    'FR-CRC':  ('완료',  ['컴플라이언스']),
    'FR-DSV':  ('완료',  ['Release Readiness(DSV 추적)']),
    'FR-KSW':  ('완료',  ['Kill Switch']),
    'FR-RBK':  ('완료',  ['Rollback(단계복구)']),
    'FR-OPD':  ('완료',  ['Telemetry', 'Ops']),
    'FR-QFL':  ('완료',  ['Telemetry(품질 피드백)']),
    'FR-EVT':  ('완료',  ['Telemetry', 'Audit']),
    'FR-VOC':  ('완료',  ['Incident']),
    'FR-RDD':  ('완료',  ['CI/CD']),
    'FR-CICD': ('완료',  ['CI/CD']),
    'FR-PDA':  ('완료',  ['Campaign(자동배포)', 'CI/CD']),
    'FR-QGV':  ('완료',  ['품질 Gate']),
    'FR-AGW':  ('완료',  ['Connector Hub']),
    'FR-LGCY': ('완료',  ['Connector Hub']),
    'FR-DSYN': ('완료',  ['Sync Logs']),
    'FR-SDVI': ('완료',  ['Connector Hub']),
    'FR-BIL':  ('완료',  ['과금 연계']),
    'FR-GLB':  ('완료',  ['글로벌 출시']),
    'FR-FSP':  ('완료',  ['현장 지원']),
    'FR-BIZ':  ('완료',  ['사업 지표']),
    'FR-PVL':  ('완료',  ['Consistency Console']),
    'FR-SPM':  ('완료',  ['API Package', 'Supplier Portal']),
    'FR-SVL':  ('완료',  ['시나리오 검증']),
    'FR-SCN':  ('완료',  ['시나리오 검증']),
}
for f in ['FR-SRT', 'FR-SDM', 'FR-CIV', 'FR-VHM', 'FR-SVS']:
    COVERAGE[f] = ('완료', ['보안 운영'])

STATUS_DESC = {
    '완료':  '플랫폼 화면에 동작 기능으로 구현됨',
    '부분':  '핵심 일부 구현·시각화, 세부 로직 일부 미완',
    '백엔드': '프론트 표시, 실제 처리는 백엔드(미연동) 필요',
    '미구현': '신규 전용 화면(placeholder)·로직 미구현',
}
FILL = {
    '완료':  PatternFill('solid', fgColor='C6EFCE'),
    '부분':  PatternFill('solid', fgColor='FFEB9C'),
    '백엔드': PatternFill('solid', fgColor='BDD7EE'),
    '미구현': PatternFill('solid', fgColor='D9D9D9'),
}
HEAD = PatternFill('solid', fgColor='0B5FFF')
HEADF = Font(color='FFFFFF', bold=True)
THIN = Border(*[Side(style='thin', color='D0D5DD')] * 4)


def cov(fam):
    return COVERAGE.get(fam, ('미구현', ['Spec Explorer']))


def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD; cell.font = HEADF
        cell.alignment = Alignment(vertical='center', horizontal='center', wrap_text=True)


def main():
    reqs = json.load(open(SPEC, encoding='utf-8')).get('requirements', [])
    # family 메타: 첫 등장 순서 보존
    fam_order, fam_meta = [], {}
    for r in reqs:
        f = r.get('family')
        if not f:
            continue
        if f not in fam_meta:
            fam_order.append(f)
            fam_meta[f] = {'sheet': r.get('sheet', ''), 'category': r.get('category', ''),
                           'name': r.get('component', ''), 'count': 0}
        fam_meta[f]['count'] += 1

    shutil.copy(SRC, OUT)
    wb = openpyxl.load_workbook(OUT)
    for name in ['플랫폼반영_요약', '플랫폼반영_Family', '플랫폼반영_FR상세']:
        if name in wb.sheetnames:
            del wb[name]

    # ── 시트1: 요약 ──
    s1 = wb.create_sheet('플랫폼반영_요약')
    s1.append(['SDV Feature Platform — 기능명세서 플랫폼 반영 현황'])
    s1['A1'].font = Font(bold=True, size=14)
    s1.append([f'총 요구사항(FR) {len(reqs)}건 · FR Family {len(fam_order)}개 · 6 기능 시트'])
    s1.append([])
    # 상태별 집계(Family / FR)
    fam_by = {k: 0 for k in STATUS_DESC}
    fr_by = {k: 0 for k in STATUS_DESC}
    for f in fam_order:
        st = cov(f)[0]; fam_by[st] += 1; fr_by[st] += fam_meta[f]['count']
    s1.append(['구현 상태', 'Family 수', 'FR 수', '비율(FR)', '설명'])
    style_header(s1, s1.max_row, 5)
    for st in ['완료', '부분', '백엔드', '미구현']:
        pct = f"{fr_by[st] / max(1, len(reqs)) * 100:.0f}%"
        s1.append([st, fam_by[st], fr_by[st], pct, STATUS_DESC[st]])
        s1.cell(row=s1.max_row, column=1).fill = FILL[st]
    s1.append(['합계', len(fam_order), len(reqs), '100%', ''])
    s1.cell(row=s1.max_row, column=1).font = Font(bold=True)
    s1.append([])
    s1.append(['※ 매핑 근거: 플랫폼 src/data/specCoverage.ts (FR family → 반영 화면·상태) + spec.json(요구사항)'])
    s1.append(['※ 상태 정의: 완료=동작 구현 / 부분=일부 구현 / 백엔드=프론트만, 처리 미연동 / 미구현=신규 화면·로직 미완'])
    for col, w in zip('ABCDE', [14, 12, 10, 10, 60]):
        s1.column_dimensions[col].width = w

    # ── 시트2: Family 단위 ──
    s2 = wb.create_sheet('플랫폼반영_Family')
    cols2 = ['기능 시트', '카테고리(개발항목)', 'FR Family', 'Family 명(컴포넌트)', 'FR 건수', '구현 상태', '반영 화면(플랫폼)']
    s2.append(cols2); style_header(s2, 1, len(cols2))
    for f in fam_order:
        m = fam_meta[f]; st, screens = cov(f)
        s2.append([m['sheet'], m['category'], f, m['name'], m['count'], st, ' · '.join(screens)])
        s2.cell(row=s2.max_row, column=6).fill = FILL[st]
    for r in range(1, s2.max_row + 1):
        for c in range(1, len(cols2) + 1):
            s2.cell(row=r, column=c).border = THIN
            s2.cell(row=r, column=c).alignment = Alignment(vertical='center', wrap_text=True)
    for col, w in zip('ABCDEFG', [22, 22, 12, 30, 9, 11, 46]):
        s2.column_dimensions[col].width = w
    s2.freeze_panes = 'A2'

    # ── 시트3: FR 단위 상세 ──
    s3 = wb.create_sheet('플랫폼반영_FR상세')
    cols3 = ['기능 시트', '카테고리', 'FR ID', '요구사항 명', 'Family', '구현 상태', '반영 화면', '적용 영역']
    s3.append(cols3); style_header(s3, 1, len(cols3))
    for r in reqs:
        f = r.get('family', ''); st, screens = cov(f)
        s3.append([r.get('sheet', ''), r.get('category', ''), r.get('id', ''), r.get('name', ''),
                   f, st, ' · '.join(screens), r.get('area', '')])
        s3.cell(row=s3.max_row, column=6).fill = FILL[st]
    for col, w in zip('ABCDEFGH', [20, 20, 16, 38, 11, 11, 40, 16]):
        s3.column_dimensions[col].width = w
    s3.freeze_panes = 'A2'

    wb.save(OUT)
    print('OK →', os.path.normpath(OUT))
    print(f'  FR={len(reqs)} families={len(fam_order)}')
    print('  Family by status:', {k: fam_by[k] for k in ['완료', '부분', '백엔드', '미구현']})
    print('  FR by status:    ', {k: fr_by[k] for k in ['완료', '부분', '백엔드', '미구현']})


if __name__ == '__main__':
    main()
